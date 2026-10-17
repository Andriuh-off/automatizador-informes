import copy
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

CARPETA = Path(__file__).parent
PLANTILLA = CARPETA / "plantilla_informe.docx"
CONTENIDO = CARPETA / "contenido.json"
CARPETA_SALIDA = CARPETA / "informes"


def limpiar(texto):
    """Quita viñetas o numeración que el modelo pudo agregar al inicio."""
    return re.sub(r"^[\s\-\•\*\u2022]+|^\d+[\.\)]\s+", "", texto).strip()


def poner_texto(parrafo, texto):
    runs = parrafo.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def buscar(doc, marcador):
    for p in doc.paragraphs:
        if p.text.strip() == marcador:
            return p
    raise ValueError(f"No se encontró el marcador {marcador} en la plantilla")


def clonar_antes(modelo, ancla, texto):
    """Copia el párrafo modelo (con su formato) y lo inserta antes del ancla."""
    nuevo = copy.deepcopy(modelo._p)
    ancla._p.addprevious(nuevo)
    poner_texto(Paragraph(nuevo, modelo._parent), texto)


def reemplazar_bloque(doc, clave, items):
    modelo = buscar(doc, "{{" + clave + "}}")
    for item in items:
        clonar_antes(modelo, modelo, limpiar(item))
    modelo._p.getparent().remove(modelo._p)


def main():
    datos = json.loads(CONTENIDO.read_text(encoding="utf-8"))
    doc = Document(PLANTILLA)

    # 1) Valores simples (datos generales)
    for p in doc.paragraphs:
        for run in p.runs:
            for campo, valor in datos["generales"].items():
                run.text = run.text.replace("{{" + campo + "}}", str(valor))

    # 2) Actividades por día (título + párrafos), se repiten por cada día
    modelo_titulo = buscar(doc, "{{ACT_TITULO}}")
    modelo_texto = buscar(doc, "{{ACT_TEXTO}}")
    for act in datos["actividades"]:
        clonar_antes(modelo_titulo, modelo_titulo, act["titulo"])
        for parrafo in act["parrafos"]:
            clonar_antes(modelo_texto, modelo_titulo, parrafo)
    modelo_titulo._p.getparent().remove(modelo_titulo._p)
    modelo_texto._p.getparent().remove(modelo_texto._p)

    # 3) Secciones generales
    for clave, items in datos["secciones"].items():
        reemplazar_bloque(doc, clave, items)

    # 4) Guardar una copia nueva en cada ejecución
    CARPETA_SALIDA.mkdir(exist_ok=True)
    salida = CARPETA_SALIDA / f"INFORME_DE_ACTIVIDADES_{datetime.now():%Y-%m-%d_%H-%M-%S}.docx"
    doc.save(salida)
    print(f"Informe generado: {salida}")


if __name__ == "__main__":
    main()